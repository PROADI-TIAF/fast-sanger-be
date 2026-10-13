import csv
import json
import urllib.request

import numpy as np
from openpyxl import load_workbook

from fast_sanger_be.cli import main
from fast_sanger_be.editors import EditorSession
from fast_sanger_be.simulate import make_example


def test_full_run_matches_simulated_truth(tmp_path):
    make_example(tmp_path)
    assert main([str(tmp_path), "--non-interactive"]) == 0
    out = tmp_path / "Final_Results"
    for path in [
        "ABE/ABE_final_results.xlsx", "ABE/Raw_Data/ABE_raw_results.xlsx", "ABE/Heatmaps/sgHBB_heatmap.png",
        "ABE/Raw_Data/Figures/sgEMX1/Dual-1_sgHBB_sgEMX1_ABE_R.png",
        "CBE/CBE_final_results.xlsx", "CBE/Raw_Data/Low_Quality/sgEMX1/CBE-2_sgEMX1_CBE_R.png", "run_log.txt",
    ]:
        assert (out / path).exists(), path

    truth = {(r["editor"], r["guide"], r["sample"], r["position"]): float(r["simulated_editing_percent"])
             for r in csv.DictReader((tmp_path / "truth.tsv").open(), delimiter="\t")}
    errors, notes = [], {}
    for editor in ("ABE", "CBE"):
        wb = load_workbook(out / editor / f"{editor}_final_results.xlsx")
        for guide in ("sgHBB", "sgEMX1"):
            if guide not in wb.sheetnames:
                continue
            rows = list(wb[guide].values)
            header = rows[0]
            for row in rows[1:]:
                notes[(editor, guide, row[0])] = (row[1], row[-1])
                for title, value in zip(header, row):
                    if title and "(%" in title and value is not None:
                        errors.append(value - truth[(editor, guide, row[0], title.split()[0])])
    assert len(errors) == 40
    assert np.mean(np.abs(errors)) < 1.5 and np.max(np.abs(errors)) < 4
    assert notes[("ABE", "sgHBB", "ABE-3")] == ("F", "R read missing")
    assert notes[("CBE", "sgEMX1", "CBE-2")][0] == "F" and "Low quality" in notes[("CBE", "sgEMX1", "CBE-2")][1]


def test_non_interactive_writes_templates(tmp_path):
    make_example(tmp_path)
    (tmp_path / "guides.txt").unlink()
    (tmp_path / "ABE-1_sgHBB_ABE_F.ab1").rename(tmp_path / "ABE-1 sgHBB fwd.ab1")
    assert main([str(tmp_path), "--non-interactive"]) == 2
    table = (tmp_path / "rename_table.tsv").read_text()
    assert "ABE-1 sgHBB fwd.ab1\tyes\t" in table
    # the guess cannot know sgHBB without a guide file: complete the row and continue
    (tmp_path / "rename_table.tsv").write_text(table.replace("\tyes\tABE-1\t\t\tF", "\tyes\tABE-1\tsgHBB\tABE\tF"))
    assert main([str(tmp_path), "--non-interactive"]) == 2
    assert (tmp_path / "ABE-1_sgHBB_ABE_F.ab1").exists()
    assert (tmp_path / "guides.txt").read_text().startswith(">")


def test_editor_session_saves_with_token(tmp_path):
    target = tmp_path / "guides.txt"
    with EditorSession("guide_editor.html", {"rows": [], "required": []}, target, open_browser=False) as session:
        page = urllib.request.urlopen(session.url).read().decode()
        assert "Guide sequences" in page and session.token in page

        def post(token):
            body = json.dumps({"token": token, "content": ">g\nACGT\n"}).encode()
            request = urllib.request.Request(session.url.rsplit("/", 1)[0] + "/save", body, method="POST")
            try:
                return urllib.request.urlopen(request).status
            except urllib.error.HTTPError as exc:
                return exc.code

        assert post("wrong") == 403 and not target.exists()
        assert post(session.token) == 200 and target.read_text() == ">g\nACGT\n"
