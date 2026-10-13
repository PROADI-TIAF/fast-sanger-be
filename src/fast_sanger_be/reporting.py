"""Excel workbooks and heatmaps for each editor (ABE, CBE)."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fast_sanger_be import __version__
from fast_sanger_be.analysis import position_label, target_positions
from fast_sanger_be.config import BASES, PRODUCT_BASE, READS, Parameters
from fast_sanger_be.naming import Guide
from fast_sanger_be.pipeline import MISSING, ReadResult, natural_key, raw_dir, summarise_guide
from fast_sanger_be.plotting import plot_heatmap

log = logging.getLogger("fast_sanger_be")

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(bold=True)
FLAG_FONT = Font(color="9C0006")


# ---------------------------------------------------------------- worksheet helpers


def _sheet_title(wb: Workbook, name: str) -> str:
    title = "".join("_" if ch in '\\/*[]:?' else ch for ch in name)[:31] or "Sheet"
    candidate, i = title, 2
    while candidate.lower() in (s.lower() for s in wb.sheetnames):
        suffix = f"_{i}"
        candidate = title[: 31 - len(suffix)] + suffix
        i += 1
    return candidate


def _write_table(ws: Worksheet, header: list[str], rows: list[list], value_columns: list[int] | None = None) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([None if isinstance(v, float) and np.isnan(v) else v for v in row])
    ws.freeze_panes = "B2"
    for idx, title in enumerate(header, start=1):
        longest = max([len(str(title))] + [len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None] or [8])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(9, longest + 2), 60)
    if value_columns and rows:
        for col in value_columns:
            letter = get_column_letter(col)
            for cell in ws[letter][1:]:
                cell.number_format = "0.0"
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{len(rows) + 1}",
                ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                               end_type="num", end_value=100, end_color="6A92C7"),
            )


def _guides_sheet(wb: Workbook, guides: list[Guide], editor: str) -> None:
    ws = wb.create_sheet("Guides")
    rows = [
        [g.name, g.sequence, len(g.sequence),
         ", ".join(position_label(g.sequence, i) for i in target_positions(g.sequence, editor))]
        for g in guides
    ]
    _write_table(ws, ["Guide", "Sequence (5'->3', protospacer)", "Length (nt)", "Target positions"], rows)


def _parameters_sheet(wb: Workbook, params: Parameters, run_info: dict[str, str]) -> None:
    ws = wb.create_sheet("Parameters")
    rows = [[k, v, ""] for k, v in run_info.items()] + [[n, v, d] for n, v, d in params.rows()]
    _write_table(ws, ["Setting", "Value", "Description"], rows)
    ws.column_dimensions["C"].width = 90


# ---------------------------------------------------------------- raw workbook


def _raw_guide_sheet(wb: Workbook, guide: Guide, editor: str, results: list[ReadResult]) -> None:
    labels = [position_label(guide.sequence, i) for i in target_positions(guide.sequence, editor)]
    by_sample: dict[str, dict[str, ReadResult]] = {}
    for r in results:
        by_sample.setdefault(r.file.sample, {})[r.file.read] = r
    rows = []
    for sample in sorted(by_sample, key=natural_key):
        for read in READS:
            r = by_sample[sample].get(read)
            if r is None:
                rows.append([sample, read, "", MISSING, ""] + [None] * len(labels))
            else:
                strand = r.hit.strand if r.hit else ""
                rows.append([sample, read, r.file.path.name, r.status_text, strand]
                            + [r.values.get(label) if r.used else None for label in labels])
    ws = wb.create_sheet(_sheet_title(wb, guide.name))
    header = ["Sample", "Read", "File", "Status", "Strand"] + [f"{label} (% {PRODUCT_BASE[editor]})" for label in labels]
    _write_table(ws, header, rows, value_columns=list(range(6, 6 + len(labels))))
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        if row[0].value != "OK":
            row[0].font = FLAG_FONT


def _qc_sheet(wb: Workbook, results: list[ReadResult]) -> None:
    rows = []
    for r in results:
        qc, q = r.qc, r.quant
        rows.append([
            r.file.path.name, r.file.sample, r.guide.name, r.file.read, r.status,
            r.hit.strand if r.hit else None, r.hit.mismatches if r.hit else None,
            qc.background_median_phred if qc else None,
            qc.background_median_peak_height if qc else None,
            100 * qc.flank_mixed_fraction if qc and qc.flank_mixed_fraction is not None else None,
            100 * qc.window_mixed_fraction if qc and qc.window_mixed_fraction is not None else None,
            *[(q.background[b] if q else None) for b in BASES],
            q.background_positions if q else None,
            "; ".join(r.reasons),
        ])
    ws = wb.create_sheet("QC")
    _write_table(ws, [
        "File", "Sample", "Guide", "Read", "Status", "Strand", "Guide mismatches",
        "Background median Phred", "Background median peak height (RFU)", "Flank overlapping peaks (%)",
        "Window overlapping peaks (%)", *[f"Background {b} (RFU)" for b in BASES],
        "Background positions", "Reasons",
    ], rows)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=16):
        for cell in row:
            cell.number_format = "0.0"


def _composition_sheet(wb: Workbook, results: list[ReadResult]) -> None:
    rows = []
    for r in results:
        if r.quant is None:
            continue
        for i, base in enumerate(r.guide.sequence):
            rows.append([r.file.path.name, r.file.sample, r.guide.name, r.file.read, r.status,
                         i + 1, base, *[float(v) for v in r.quant.percentages[i]]])
    ws = wb.create_sheet("Base_composition")
    _write_table(ws, ["File", "Sample", "Guide", "Read", "Status", "Guide position", "Reference base",
                      "% A", "% C", "% G", "% T"], rows)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=11):
        for cell in row:
            cell.number_format = "0.00"


def write_raw_workbook(path: Path, editor: str, guides: list[Guide], results: list[ReadResult],
                       params: Parameters, run_info: dict[str, str]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for guide in guides:
        _raw_guide_sheet(wb, guide, editor, [r for r in results if r.guide.name == guide.name])
    _qc_sheet(wb, results)
    _composition_sheet(wb, results)
    _guides_sheet(wb, guides, editor)
    _parameters_sheet(wb, params, run_info)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- final workbook


def write_final_workbook(path: Path, editor: str, guides: list[Guide], results: list[ReadResult],
                         params: Parameters, run_info: dict[str, str]) -> dict[str, list]:
    wb = Workbook()
    wb.remove(wb.active)
    summaries_by_guide = {}
    for guide in guides:
        labels = [position_label(guide.sequence, i) for i in target_positions(guide.sequence, editor)]
        summaries = summarise_guide(results, guide, editor, params)
        summaries_by_guide[guide.name] = summaries
        rows = [
            [s.sample, s.reads_used] + [s.values.get(label) for label in labels]
            + [s.max_difference, "; ".join(s.notes)]
            for s in summaries
        ]
        ws = wb.create_sheet(_sheet_title(wb, guide.name))
        header = (["Sample", "Reads used"] + [f"{label} (% {PRODUCT_BASE[editor]})" for label in labels]
                  + ["Max F/R difference (pp)", "Notes"])
        _write_table(ws, header, rows, value_columns=list(range(3, 3 + len(labels))))
        diff_letter = get_column_letter(3 + len(labels))
        for cell in ws[diff_letter][1:]:
            cell.number_format = "0.0"
        for row in ws.iter_rows(min_row=2, min_col=len(header), max_col=len(header)):
            if row[0].value:
                row[0].font = FLAG_FONT

    excluded = [
        [r.file.path.name, r.file.sample, r.guide.name, r.file.read, r.status, "; ".join(r.reasons)]
        for r in results if not r.used
    ]
    present = {(r.guide.name, r.file.sample, r.file.read) for r in results}
    samples_per_guide = {(r.guide.name, r.file.sample) for r in results}
    for guide_name, sample in sorted(samples_per_guide):
        for read in READS:
            if (guide_name, sample, read) not in present:
                excluded.append(["", sample, guide_name, read, MISSING, "no .ab1 file for this read"])
    ws = wb.create_sheet("Excluded_reads")
    _write_table(ws, ["File", "Sample", "Guide", "Read", "Status", "Reason"], excluded)

    _guides_sheet(wb, guides, editor)
    _parameters_sheet(wb, params, run_info)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return summaries_by_guide


def write_editor_outputs(out_root: Path, editor: str, guides: list[Guide], results: list[ReadResult],
                         params: Parameters, run_info: dict[str, str]) -> dict[str, list]:
    """Write the raw and final workbooks and the heatmaps. Returns the sample summaries per guide."""
    editor_dir = out_root / editor
    write_raw_workbook(raw_dir(out_root, editor) / f"{editor}_raw_results.xlsx",
                       editor, guides, results, params, run_info)
    summaries = write_final_workbook(editor_dir / f"{editor}_final_results.xlsx",
                                     editor, guides, results, params, run_info)
    for guide in guides:
        rows = summaries[guide.name]
        if not rows:
            continue
        plot_heatmap(
            guide.name, guide.sequence, [s.sample for s in rows],
            {s.sample: s.values for s in rows}, editor,
            editor_dir / "Heatmaps" / f"{guide.name}_heatmap.{params.figure_format}", params,
        )
    return summaries


def run_info(input_dir: Path | str, n_files: int) -> dict[str, str]:
    return {
        "Software": f"Fast-Sanger-BE {__version__}",
        "Run date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Input folder": str(input_dir),
        "Number of .ab1 files": str(n_files),
    }
